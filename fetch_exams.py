import os
import json
import re
import requests
from io import BytesIO
import openpyxl

EXAM_URL = "https://docs.google.com/spreadsheets/d/1x6k9Ub61Fg3iIUFNRoOl3J9omc_Q4Dbp/export?format=xlsx"
OUTPUT_FILE = "frontend/public/exams_s1.json"

def fetch_and_parse_exams():
    print(f"Downloading Exam Sheet: {EXAM_URL}")
    resp = requests.get(EXAM_URL)
    if resp.status_code != 200:
        print(f"Failed to download sheet. Status Code: {resp.status_code}. Ensure it is public.")
        return None

    wb = openpyxl.load_workbook(filename=BytesIO(resp.content), data_only=True)
    
    is_final_draft = False
    exams = []
    
    dept_patterns = {
        'CS': r'\b(CS|BCS)\b',
        'SE': r'\b(SE|BSE)\b',
        'AI': r'\b(AI|BAI)\b',
        'DS': r'\b(DS|BDS)\b',
        'CY': r'\b(CY|BCY)\b',
        'AF': r'\b(AF|BAF)\b',
        'FT': r'\b(FT|BFT)\b',
        'BBA': r'\bBBA\b',
        'BA': r'\b(BA|BSBA)\b',
        'BCE': r'\bBCE\b',
        'BEE': r'\bBEE\b'
    }

    for sheet_name in ['FSC', 'FSM', 'FSE']:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        
        # 1. State Tracking: Check if document is "Final"
        for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
            for cell_val in row:
                if cell_val and isinstance(cell_val, str) and "final" in cell_val.lower():
                    is_final_draft = True
                    break

        # 2. Get Time Slots from Row 4
        time_slots = {}
        for idx, cell in enumerate(ws[4]):
            if cell.value and "to" in str(cell.value):
                time_slots[idx] = str(cell.value).strip()

        current_date = "Unknown Date"
        
        # 3. Iterate over rows starting from 5
        for row in ws.iter_rows(min_row=5):
            # Track the date
            date_val = row[0].value
            if date_val:
                current_date = str(date_val).split(' ')[0]

            # Check each time slot column
            for col_idx, time_str in time_slots.items():
                cell = row[col_idx]
                if not cell.value:
                    continue
                    
                raw_text = str(cell.value).strip()
                lines = [line.strip() for line in raw_text.split('\n') if line.strip()]
                
                if not lines:
                    continue
                    
                course_name = lines[0]
                if len(lines) > 1 and not re.search(r'(BS|MS|PhD|BBA|BAF|BCE|BEE|MB|MEE|PEE)\b', lines[1].upper()) and not re.search(r'\b202\d\b', lines[1]):
                    course_name = f"{lines[0]} {lines[1]}"
                
                # Extract batch
                batch = "Unknown"
                for line in lines:
                    match = re.search(r'\b(202\d)\b', line)
                    if match:
                        batch = match.group(1)
                        break
                        
                # Exclude MS/PhD
                has_ms_phd = any(re.search(r'\b(MS|PHD|MB|MEE|PEE)\b|MSBA', line, re.IGNORECASE) for line in lines)
                if has_ms_phd:
                    continue 

                # Aggressive Department Extraction
                depts = set()
                for d, pattern in dept_patterns.items():
                    if re.search(pattern, raw_text, re.IGNORECASE):
                        depts.add(d)

                # If no departments matched, this is either a signature (Abdul Hameed) or a copied header row, so skip.
                if not depts:
                    continue

                # 3. Color Detection for "Repeated" courses (Dark Red)
                is_repeat = False
                if cell.fill and cell.fill.start_color:
                    color_hex = str(cell.fill.start_color.rgb)
                    if color_hex.startswith("FF") and color_hex[2:4] in ["FF", "C0", "80"] and color_hex[4:] == "0000":
                        is_repeat = True

                exams.append({
                    "id": f"EXAM-{course_name.replace(' ', '')}-{batch}",
                    "departments": list(depts),
                    "batch": batch,
                    "course_name": course_name,
                    "date": current_date,
                    "time": time_str,
                    "is_repeat": is_repeat
                })

    return {
        "is_final_draft": is_final_draft,
        "exams": exams
    }

def diff_and_save(new_data):
    if os.path.exists(OUTPUT_FILE):
        try:
            with open(OUTPUT_FILE, "r") as f:
                old_data = json.load(f)
            if json.dumps(old_data, sort_keys=True) == json.dumps(new_data, sort_keys=True):
                print("No changes detected in Exam Schedule.")
                return False
        except Exception as e:
            print(f"Error reading old data: {e}")
            
    with open(OUTPUT_FILE, "w") as f:
        json.dump(new_data, f, indent=2)
    print(f"Changes detected! Saved new exam schedule to {OUTPUT_FILE}")
    return True

if __name__ == "__main__":
    data = fetch_and_parse_exams()
    if data:
        changed = diff_and_save(data)
        if changed:
            exit(1)
        else:
            exit(0)
